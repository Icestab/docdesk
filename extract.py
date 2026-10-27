"""文档 -> 页块列表。只读，不写任何磁盘文件。

每个 "page" 是 {"label": "第3页", "text": "..."}，label 会出现在引用里。
"""

from __future__ import annotations

import html
import io
import os
import re
import shutil
import subprocess
import xml.etree.ElementTree as ET

MAX_BYTES = 120 * 1024 * 1024
DPI = 300
OCR_LANGS = "chi_sim+eng"

# 优先用系统装的 tesseract；没有就用本地解压的那份（无需 root）
TESS_LOCAL = os.environ.get("DOCDESK_TESS", os.path.expanduser("~/opt/tess"))


def _find_tesseract():
    """优先系统安装（且带中文包），否则用本地解压的那份。"""
    env = dict(os.environ)
    data_dirs = ["/usr/share/tessdata", "/usr/share/tesseract-ocr/5/tessdata",
                 "/usr/share/tesseract-ocr/4.00/tessdata"]
    system = shutil.which("tesseract")
    if system:
        for d in data_dirs:
            if os.path.exists(os.path.join(d, "chi_sim.traineddata")):
                env["TESSDATA_PREFIX"] = d
                return system, env
        if os.environ.get("TESSDATA_PREFIX"):
            return system, env

    binary = os.path.join(TESS_LOCAL, "usr/bin/tesseract")
    data = os.path.join(TESS_LOCAL, "tessdata")
    if os.path.exists(binary) and os.path.isdir(data):
        env["LD_LIBRARY_PATH"] = os.path.join(TESS_LOCAL, "usr/lib")
        env["TESSDATA_PREFIX"] = data
        return binary, env
    return (system, env) if system else (None, None)


TESS_BIN, TESS_ENV = _find_tesseract()

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


class ExtractError(Exception):
    pass


def _clean(text: str) -> str:
    text = text.replace("\x00", " ").replace("\r\n", "\n").replace("\r", "\n")
    # 英文断词连字符
    text = re.sub(r"(?<=[A-Za-z])-\n(?=[a-z])", "", text)
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def ocr_ready() -> bool:
    return TESS_BIN is not None


def _despace_cjk(text: str) -> str:
    """tesseract 的 chi_sim 习惯在每个汉字之间插空格，去掉。"""
    text = re.sub(r"(?<=[\u4e00-\u9fff])\s+(?=[\u4e00-\u9fff])", "", text)
    text = re.sub(r"(?<=[\u4e00-\u9fff])\s+(?=[，。；：、！？%）】」』])", "", text)
    text = re.sub(r"(?<=[（【「『])\s+", "", text)
    text = re.sub(r"\s+(?=[（【「『])", "", text)
    return text


def _ocr_png(png: bytes) -> str:
    proc = subprocess.run(
        [TESS_BIN, "stdin", "stdout", "-l", OCR_LANGS, "--psm", "3",
         "-c", "preserve_interword_spaces=1"],
        input=png, capture_output=True, env=TESS_ENV, timeout=240)
    if proc.returncode != 0:
        raise ExtractError("OCR 失败：" + proc.stderr.decode("utf-8", "replace")[:160])
    return _despace_cjk(proc.stdout.decode("utf-8", "replace")).strip()


def _page_png(page, dpi: int = DPI) -> bytes:
    image = page.render(scale=dpi / 72).to_pil().convert("RGB")
    buf = io.BytesIO()
    image.save(buf, format="PNG")
    return buf.getvalue()


def extract(name: str, data: bytes) -> dict:
    """返回 {"pages": [...], "kind": str, "scanned": [页号...]}"""
    if len(data) > MAX_BYTES:
        raise ExtractError(f"文件太大（{len(data)/1e6:.0f} MB），上限 120 MB")
    ext = name.rsplit(".", 1)[-1].lower() if "." in name else ""

    if ext == "pdf":
        pages, scanned, ocr_used = _pdf(data)
        kind = "pdf"
    elif ext in ("docx", "docm"):
        pages, scanned, ocr_used = _docx(data), [], []
        kind = "docx"
    elif ext in ("xlsx", "xlsm"):
        pages, scanned, ocr_used = _xlsx(data), [], []
        kind = "xlsx"
    elif ext in ("html", "htm"):
        pages, scanned, ocr_used = _html(data), [], []
        kind = "html"
    elif ext in ("png", "jpg", "jpeg", "webp", "bmp", "tif", "tiff"):
        pages = _img(data)
        scanned, ocr_used = [], [1]
        kind = "image"
    elif ext in ("txt", "md", "markdown", "csv", "tsv", "log", "json",
                 "yaml", "yml", "srt", "vtt", "tex", "rst", "ini", ""):
        pages, scanned, ocr_used = _text(data), [], []
        kind = "text"
    elif ext in ("doc", "xls", "ppt", "pptx", "odt"):
        raise ExtractError(f"暂不支持 .{ext}，请先另存为 PDF 或 docx")
    else:
        raise ExtractError(f"不认识的格式 .{ext}")

    total = sum(len(p["text"]) for p in pages)
    if total < 20:
        hint = ("扫描件需要 OCR，但本机没装 tesseract"
                if not ocr_ready() else "这些页 OCR 也没识别出文字")
        raise ExtractError("没有抽出可读文字：" + hint)
    return {"pages": pages, "kind": kind, "chars": total,
            "scanned": scanned, "ocr": ocr_used}


# ---------------------------------------------------------------- PDF

def _pdf(data: bytes):
    import pypdfium2 as pdfium

    try:
        doc = pdfium.PdfDocument(data)
    except Exception as exc:  # pragma: no cover
        raise ExtractError(f"PDF 打不开：{exc}") from exc

    pages, scanned, ocr_used = [], [], []
    try:
        total = len(doc)
        for i in range(total):
            page = doc[i]
            text = _clean(page.get_textpage().get_text_bounded())
            if len(text) < 30 and ocr_ready():
                try:
                    text = _ocr_png(_page_png(page))
                    if len(text) >= 10:
                        ocr_used.append(i + 1)
                    else:
                        text = ""
                except Exception:
                    text = ""
            page.close()
            if len(text) < 30:
                scanned.append(i + 1)
            if text:
                mark = "（OCR）" if (i + 1) in ocr_used else ""
                pages.append({"label": f"第{i+1}页{mark}", "text": text})
    finally:
        doc.close()

    if not pages:
        if not ocr_ready():
            raise ExtractError(
                f"这 {len(scanned)} 页都没有文字层，是扫描件；本机没装 tesseract，识别不了。")
        raise ExtractError("所有页 OCR 之后仍然没有文字，可能是空白页或拍照太糊。")
    return pages, scanned, ocr_used


def _img(data: bytes):
    if not ocr_ready():
        raise ExtractError("图片要靠 OCR 识别，但本机没装 tesseract")
    text = _ocr_png(data)
    if len(text) < 10:
        raise ExtractError("这张图没识别出文字（太糊、太斜或者没有字）")
    return [{"label": "图片（OCR）", "text": text}]


# ---------------------------------------------------------------- docx

def _docx(data: bytes):
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(io.BytesIO(data))
    blocks: list[str] = []
    body = doc.element.body
    for child in body.iterchildren():
        if child.tag == W + "p":
            t = Paragraph(child, doc).text.strip()
            if t:
                blocks.append(t)
        elif child.tag == W + "tbl":
            rows = []
            for row in Table(child, doc).rows:
                cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                if any(cells):
                    rows.append(" | ".join(cells))
            if rows:
                blocks.append("\n".join(rows))

    # 按字符量切页，方便引用定位
    pages, buf, size, n = [], [], 0, 1
    for b in blocks:
        buf.append(b)
        size += len(b)
        if size >= 1800:
            pages.append({"label": f"第{n}段", "text": "\n\n".join(buf)})
            n += 1
            buf, size = [], 0
    if buf:
        pages.append({"label": f"第{n}段", "text": "\n\n".join(buf)})
    return pages


# ---------------------------------------------------------------- xlsx

def _xlsx(data: bytes):
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    pages = []
    try:
        for ws in wb.worksheets:
            lines, n = [], 0
            for row in ws.iter_rows(values_only=True):
                cells = ["" if v is None else str(v) for v in row]
                while cells and not cells[-1]:
                    cells.pop()
                if not cells:
                    continue
                lines.append("\t".join(cells))
                if len(lines) >= 120:
                    n += 1
                    pages.append({"label": f"{ws.title}·{n}", "text": "\n".join(lines)})
                    lines = []
            if lines:
                n += 1
                pages.append({"label": f"{ws.title}·{n}", "text": "\n".join(lines)})
    finally:
        wb.close()
    return pages


# ---------------------------------------------------------------- text/html

def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "gb18030", "big5", "utf-16"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _text(data: bytes):
    text = _clean(_decode(data))
    pages, step = [], 4000
    for i in range(0, len(text), step):
        pages.append({"label": f"§{i//step+1}", "text": text[i:i + step]})
    return pages or [{"label": "§1", "text": text}]


def _html(data: bytes):
    raw = _decode(data)
    raw = re.sub(r"(?is)<(script|style|head)[^>]*>.*?</\1>", " ", raw)
    raw = re.sub(r"(?i)<br\s*/?>|</p>|</div>|</tr>|</h[1-6]>", "\n", raw)
    raw = re.sub(r"(?s)<[^>]+>", " ", raw)
    text = _clean(html.unescape(raw))
    text = re.sub(r"[ \t]{2,}", " ", text)
    return _text(text.encode())
